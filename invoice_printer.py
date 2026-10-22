"""
Invoice Printer  (v2)
---------------------
Search an invoice number -> prints ONLY the INV and PL sheets.

New in v2
  * Printer selection        (dropdown; the Windows default printer is restored afterwards)
  * Batch printing           (paste many invoice numbers, print them all in one go)
  * Auto-print watch folder  (new Excel files that land in a folder are printed by themselves)
  * Print log                (print_log.csv + a log viewer window)
  * Conditional print        (INV-1 / PL-1 / INV-2 / PL-2 are printed ONLY when the sheet has line data)
  * No duplicate printing    (INV + INV, INV + INV-1 ... with the same invoice no. and content -> printed once)

Requirements : Windows + Microsoft Excel (desktop) installed
Install once : pip install pywin32 openpyxl
Run          : double-click Invoice_Printer.bat   (or: python invoice_printer.py)
"""
import csv
import hashlib
import json
import os
import queue
import re
import sys
import threading
from datetime import datetime
from pathlib import Path

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except Exception:  # lets the non-GUI logic be imported/tested without tkinter
    tk = None

APP_VERSION = "2.2.0"                    # single source of truth - the installer reads this too
APP_NAME = "Invoice Printer"
FROZEN = getattr(sys, "frozen", False)   # True when running as the packaged .exe


def _resolve_app_dir():
    """Folder for config / log / state.
    Script  -> folder of this .py file.
    .exe    -> folder of the .exe (portable). If that folder is read-only
               (e.g. C:/Program Files) -> %APPDATA%/InvoicePrinter instead.
    (Do NOT use __file__ in a one-file .exe: it points to a temp folder that is wiped on exit.)"""
    base = Path(sys.executable).resolve().parent if FROZEN else Path(__file__).resolve().parent
    try:
        probe = base / ".write_test"
        probe.write_text("x")
        probe.unlink()
        return base
    except Exception:
        alt = Path(os.environ.get("APPDATA", str(Path.home()))) / "InvoicePrinter"
        alt.mkdir(parents=True, exist_ok=True)
        return alt


APP_DIR = _resolve_app_dir()
CONFIG_PATH = APP_DIR / "invoice_printer_config.json"
LOG_PATH = APP_DIR / "print_log.csv"
STATE_PATH = APP_DIR / "auto_print_state.json"
ERROR_LOG_PATH = APP_DIR / "error_log.txt"

DEFAULT_PRINTER_LABEL = "(Windows default)"

DEFAULT_CONFIG = {
    "root_folder": "",                    # top folder that contains the month folders
    "target_folder_name": "inv&pl",       # only search inside folders with this name ("" = everywhere)
    "extensions": [".xlsx", ".xlsm", ".xls"],
    "sheet_include": ["^INV", "^PL"],     # regex, case-insensitive: sheets to print
    "sheet_exclude": ["BKK", "^ERP"],     # regex, case-insensitive: never print these
    "copies": 1,
    "confirm_before_print": True,
    # --- new in v2 ---
    "printer": "",                        # "" = Windows default printer
    "watch_folder": "",                   # folder watched for NEW Excel files ("" = root_folder)
    "watch_name_regex": "",               # only auto-print file names matching this regex ("" = all)
    "watch_interval_sec": 5,              # how often the folder is checked
    "auto_print_enabled": False,          # start auto-printing when the app opens
    # --- conditional print ---
    # Sheets matching these regexes (case-insensitive) are printed ONLY when they hold line data
    # (a quantity > 0 under the QUANTITY column). Empty / unused split sheets are skipped.
    # Set to [] to turn the rule off.
    "conditional_sheets": ["^INV-1", "^PL-1", "^INV-2", "^PL-2"],
    # Two INV sheets (or two PL sheets) with the SAME invoice number and content - e.g. INV and INV, or
    # INV and INV-1 that still carries the same INV.NO. - are printed only once (first one is kept).
    "skip_duplicate_sheets": True,
    # Excel names a copied sheet "INV(KNITTED) (2)", "INV (3)" ...  These are NOT variants: only the first sheet
    # (INV(KNITTED)) is printed.  Real variants INV-1 / PL-1 / INV-2 ... are separate sheets and still print
    # ("All sheets" box, or typing C264416-1).   False = print every "(n)" sheet too.
    "skip_numbered_copies": True,
    # True = if Excel cannot be switched to the printer chosen in the window, print NOTHING and report it
    # (never print on the wrong printer).  False = warn in the log and print on whatever Excel is using.
    "strict_printer": True,
    # --- PDF mode (checkbox "Print PDF files"): search *.pdf in the root folder and print the PDF files as they are
    "print_pdf_files": False,
    "pdf_extensions": [".pdf"],
    "pdf_target_folder_name": "",        # "" = search every sub-folder (the inv&pl folder rule is for Excel files)
    "pdf_sumatra_silent": False,         # True = hide SumatraPDF's own error messages (False = show them, easier to diagnose)
    "pdf_sumatra_extra": "",             # extra SumatraPDF -print-settings, e.g. "fit" or "fit,monochrome" ("" = none)
    "pdf_viewer_path": "",               # optional: full path of SumatraPDF.exe / AcroRd32.exe; "" = auto-detect
    # True = all chosen sheets of one workbook (INV + PL ...) go to the printer as ONE print job, so a printer that
    # asks for a code / PIN / password per job asks once per file instead of once per sheet.
    # (With copies > 1 the pages then come out as complete sets: INV,PL,INV,PL.)  False = one job per sheet.
    "print_one_job_per_file": True,
    # --- PDF quality ---
    # Fonts that are mis-spelled in the workbooks (e.g. "Time New Roman") are swapped on a temporary copy
    # before the PDF is made, so Excel does not substitute a wider font (stretched / clipped text).
    "font_fixes": {"Time New Roman": "Times New Roman", "Times New Roman ": "Times New Roman"},
    # Printer whose metrics Excel uses to lay out the PDF. "" = chosen printer, else Windows default,
    # else "Microsoft Print to PDF". Use the printer you really print on for a PDF identical to paper.
    "pdf_layout_printer": "",
    # How the PDF is made. "printer" = print through "Microsoft Print to PDF" (same engine as real paper, avoids the
    # stretched / letter-spaced text of Excel's own PDF export). "export" = Excel's ExportAsFixedFormat.
    # "printer" falls back to "export" automatically if that printer is missing or fails.
    "pdf_engine": "printer",
    "all_variants": False,                # checkbox "All sheets": print INV, INV-1 ... INV-n without typing -1, -2
}

JOB_LOCK = threading.Lock()   # only one Excel job at a time (manual, batch or auto)
_LOG_LOCK = threading.Lock()


# --------------------------------------------------------------------------
# Config / state
# --------------------------------------------------------------------------
def load_config():
    cfg = dict(DEFAULT_CONFIG)
    try:
        cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
    except Exception:
        pass
    return cfg


def save_config(cfg):
    CONFIG_PATH.write_text(json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8")


def load_state():
    try:
        return json.loads(STATE_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {"hashes": []}


def save_state(state):
    try:
        STATE_PATH.write_text(json.dumps(state, indent=2), encoding="utf-8")
    except Exception:
        pass


# --------------------------------------------------------------------------
# Print log (CSV)
# --------------------------------------------------------------------------
LOG_FIELDS = ["time", "mode", "file", "sheets", "copies", "printer", "status", "message"]


def write_log(mode, file, sheets, copies, printer, status, message=""):
    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        mode,
        file,
        " | ".join(sheets) if sheets else "",
        copies,
        printer or "(default)",
        status,
        message.replace("\n", " "),
    ]
    try:
        with _LOG_LOCK:
            is_new = not LOG_PATH.exists()
            with open(LOG_PATH, "a", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                if is_new:
                    w.writerow(LOG_FIELDS)
                w.writerow(row)
    except Exception:
        pass  # never let logging break printing
    return row


def read_log(limit=300):
    """Newest rows first."""
    if not LOG_PATH.exists():
        return []
    try:
        with _LOG_LOCK, open(LOG_PATH, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))[1:]
    except Exception:
        return []
    return rows[-limit:][::-1]


def _read_log_all():
    """Every log row, oldest first (header removed)."""
    if not LOG_PATH.exists():
        return []
    try:
        with _LOG_LOCK, open(LOG_PATH, newline="", encoding="utf-8-sig") as f:
            return list(csv.reader(f))[1:]
    except Exception:
        return []


def rewrite_log(keep_rows):
    """Replace print_log.csv by header + keep_rows (oldest first).
    The previous file is first copied to print_log_backup.csv (one backup, overwritten each time)."""
    import shutil
    with _LOG_LOCK:
        if LOG_PATH.exists():
            try:
                shutil.copyfile(LOG_PATH, LOG_PATH.with_name("print_log_backup.csv"))
            except Exception:
                pass
        with open(LOG_PATH, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(LOG_FIELDS)
            w.writerows(keep_rows)


def clear_log(older_than_days=None):
    """Clear the whole log, or only rows older than N days. Returns how many rows were removed."""
    from datetime import timedelta
    rows = _read_log_all()
    if older_than_days is None:
        keep = []
    else:
        limit = datetime.now() - timedelta(days=int(older_than_days))
        keep = []
        for r in rows:
            try:
                if datetime.strptime(r[0], "%Y-%m-%d %H:%M:%S") >= limit:
                    keep.append(r)
            except Exception:
                keep.append(r)          # unreadable date -> keep it
    rewrite_log(keep)
    return len(rows) - len(keep)


def delete_log_rows(rows_to_delete):
    """Remove the given rows (lists of strings, as shown in the viewer). Returns how many were removed."""
    remove = [list(r) for r in rows_to_delete]
    keep, removed = [], 0
    for r in _read_log_all():
        padded = (r + [""] * len(LOG_FIELDS))[:len(LOG_FIELDS)]
        if padded in remove:
            remove.remove(padded)
            removed += 1
        else:
            keep.append(r)
    rewrite_log(keep)
    return removed


# --------------------------------------------------------------------------
# Printers
# --------------------------------------------------------------------------
def list_printers():
    try:
        import win32print
        flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
        return sorted({p[2] for p in win32print.EnumPrinters(flags)})
    except Exception:
        return []


def _switch_default_printer(name):
    """Make `name` the Windows default printer. Returns the previous default (or None if unchanged)."""
    import win32print
    old = win32print.GetDefaultPrinter()
    if old == name:
        return None
    win32print.SetDefaultPrinter(name)
    return old


def excel_installed():
    """True if desktop Excel is registered (COM class Excel.Application)."""
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Excel.Application"))
        return True
    except Exception:
        return sys.platform != "win32"   # not Windows (dev/testing) -> don't nag


# --------------------------------------------------------------------------
# Searching
# --------------------------------------------------------------------------
def _norm(text):
    return re.sub(r"\s+", "", text.lower())


def find_invoices(root, query, target_folder="", extensions=(".xlsx", ".xlsm", ".xls"), cancel=None,
                  base_fallback=True):
    """Search all sub-folders of `root` for Excel files whose name contains `query`."""
    query = query.strip().lower()
    if not query or not root or not os.path.isdir(root):
        return []
    base, variant = parse_query(query)
    target = _norm(target_folder)
    exts = tuple(e.lower() for e in extensions)
    hits = []
    for dirpath, _dirs, files in os.walk(root):
        if cancel is not None and cancel.is_set():
            break
        if target and target not in _norm(os.path.basename(dirpath)):
            continue  # keep walking down, but ignore files outside the target folder
        for name in files:
            if name.startswith("~$") or not name.lower().endswith(exts):
                continue  # skip Excel lock/temp files and non-Excel files
            if query in name.lower():
                full = os.path.join(dirpath, name)
                try:
                    mtime = os.path.getmtime(full)
                except OSError:
                    mtime = 0
                hits.append({
                    "path": full,
                    "name": name,
                    "folder": os.path.relpath(dirpath, root),
                    "mtime": mtime,
                    "variant": variant,
                })
    hits.sort(key=lambda h: h["mtime"], reverse=True)
    if not hits and variant and base_fallback:
        # 'C264958-1' usually lives in the file 'C264958-....xlsx' -> search by the base number
        hits = find_invoices(root, base, target_folder, extensions, cancel)
        for h in hits:
            h["variant"] = variant
    return hits


def find_many(root, numbers, target_folder="", extensions=(".xlsx", ".xlsm", ".xls"), cancel=None,
              base_fallback=True):
    """ONE pass over the folder tree for many invoice numbers.
    Returns {number: [hits]} - matches inside the target folder come first, then newest first."""
    result = {n: [] for n in numbers}
    if not numbers or not root or not os.path.isdir(root):
        return result
    target = _norm(target_folder)
    exts = tuple(e.lower() for e in extensions)
    lowered = [(n, n.lower()) for n in numbers]
    for dirpath, _dirs, files in os.walk(root):
        if cancel is not None and cancel.is_set():
            break
        in_target = (not target) or (target in _norm(os.path.basename(dirpath)))
        for name in files:
            if name.startswith("~$") or not name.lower().endswith(exts):
                continue
            low = name.lower()
            for n, nl in lowered:
                if nl in low:
                    full = os.path.join(dirpath, name)
                    try:
                        mtime = os.path.getmtime(full)
                    except OSError:
                        mtime = 0
                    result[n].append({
                        "path": full,
                        "name": name,
                        "folder": os.path.relpath(dirpath, root),
                        "mtime": mtime,
                        "in_target": in_target,
                        "variant": parse_query(n)[1],
                    })
    for n in result:
        result[n].sort(key=lambda h: (not h["in_target"], -h["mtime"]))
    # 'C264958-1' not found by its full name -> look again with the base number 'C264958'
    retry = {}
    for n in numbers:
        base, variant = parse_query(n)
        if not result[n] and variant and base_fallback:
            retry.setdefault(base, []).append(n)
    if retry:
        again = find_many(root, list(retry), target_folder, extensions, cancel)
        for base, ns in retry.items():
            for n in ns:
                result[n] = [dict(h, variant=parse_query(n)[1]) for h in again[base]]
    return result


def parse_numbers(text):
    """'INV001, INV002\nINV003' -> ['INV001', 'INV002', 'INV003'] (no duplicates, order kept)."""
    out, seen = [], set()
    base = None
    for part in re.split(r"[,;\s]+", text.strip()):
        if re.fullmatch(r"-(\d+|\*)", part or "") and base:
            part = base + part              # 'C264958,-1,-2' -> C264958, C264958-1, C264958-2
        elif part:
            base = parse_query(part)[0] if parse_query(part)[1] is not None else part
        if part and part.lower() not in seen:
            seen.add(part.lower())
            out.append(part)
    return out


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_numbers_from_table(rows):
    """rows = list of row lists (Excel / CSV).  Finds the invoice column and returns invoice numbers.
    Header 'Invoice' / 'Invoice no' / 'INV' picks the column; optional 'Variant' / 'Suffix' column
    (1, 2 ...) turns C264958 + 1 into C264958-1.  Without a header the first column is used.
    A cell may also hold 'C264958,-1,-2'."""
    head_re = re.compile(r"inv|number|no\.?$|order", re.I)
    var_re = re.compile(r"variant|suffix|sub|part", re.I)
    inv_col = var_col = None
    start = 0
    for ri, row in enumerate(rows[:6]):
        cells = [_cell_text(c) for c in row]
        for ci, c in enumerate(cells):
            if c and inv_col is None and head_re.search(c) and not _QUERY_RE.match(c):
                inv_col, start = ci, ri + 1
            elif c and var_col is None and var_re.search(c) and not _QUERY_RE.match(c):
                var_col = ci
        if inv_col is not None:
            break
    out = []
    for row in rows[start:]:
        cells = [_cell_text(c) for c in row]
        if inv_col is None:                      # no header: first non-empty column of the row
            val = next((c for c in cells if c), "")
            if not val or not (_QUERY_RE.match(val) or re.fullmatch(r"-\d+", val) or "," in val):
                continue                          # skips titles like 'Invoice list'
        else:
            val = cells[inv_col] if inv_col < len(cells) else ""
            if not val:
                continue
            if var_col is not None and var_col < len(cells) and cells[var_col]:
                for v in re.split(r"[,;\s]+", cells[var_col]):
                    v = v.strip().lstrip("-")
                    if v.lower() in ("all", "*"):
                        v = "*"
                    if v:
                        out.append(f"{val}-{v}")
                continue
        out.append(val)
    return parse_numbers(", ".join(out))


def read_numbers_from_file(path):
    ext = Path(path).suffix.lower()
    if ext in (".csv", ".txt"):
        import csv
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        raise ValueError("Please use an .xlsx, .xlsm or .csv file.")
    return read_numbers_from_table(rows)


def save_template(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(["Invoice no.", "Variant (optional)"])
    for row in (["C264958", ""], ["C264958", "1"], ["C264958", "2"], ["C264959", ""], ["C264960-1", ""]):
        ws.append(row)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    wb.save(path)


def pick_sheets(names, include, exclude):
    inc = [re.compile(p, re.I) for p in include]
    exc = [re.compile(p, re.I) for p in exclude]
    return [n for n in names
            if not any(r.search(n) for r in exc) and any(r.search(n) for r in inc)]


# --- invoice-number variants:  C264958 -> only INV + PL      C264958-1 -> only INV-1 + PL-1
ALL_VARIANTS = "*"
_QUERY_RE = re.compile(r"^\s*([A-Za-z]*\d{3,})(?:\s*-\s*(\d+|\*))?\s*$")   # C264958-*  = ALL INV/PL sheets
_SHEET_VARIANT_RE = re.compile(r"^\s*(?:INV|PL)\s*-\s*(\d+)", re.I)


def parse_query(query):
    """'C264958' -> ('C264958', '')   'C264958-1' -> ('C264958', '1')
    Anything that does not look like an invoice number (e.g. 'korea') -> (query, None) = no sheet filter."""
    m = _QUERY_RE.match(query or "")
    if not m:
        return (query or "").strip(), None
    return m.group(1), (m.group(2) or "")


def sheet_variant(name):
    """'INV(WOVEN)' / 'PL(WOVEN) (2)' -> ''     'INV-1(WOVEN) (2)' / 'PL-1(WOVEN) (3)' -> '1'"""
    m = _SHEET_VARIANT_RE.match(name)
    return m.group(1) if m else ""


def filter_by_variant(picked, variant):
    """variant None -> unchanged.  '' -> only INV / PL.  '1' -> only INV-1 / PL-1.
    Strict: C264958-3 with no INV-3 / PL-3 sheet gives an EMPTY list (nothing is printed)."""
    if variant is None or variant == ALL_VARIANTS:
        return list(picked)
    wanted = {variant} if isinstance(variant, str) else set(variant)
    return [n for n in picked if sheet_variant(n) in wanted]


def parse_multi(query):
    """'C264958,-1,-2' -> ('C264958', ('', '1', '2')).  Not a multi-number query -> None."""
    nums = parse_numbers(query or "")
    if len(nums) < 2:
        return None
    parsed = [parse_query(n) for n in nums]
    if any(v is None for _b, v in parsed) or len({b.lower() for b, _v in parsed}) != 1:
        return None
    return parsed[0][0], tuple(v for _b, v in parsed)


def with_all_variants(numbers):
    """'Print every INV / PL sheet without typing -1, -2 ...':  C264958 -> C264958-*.
    Numbers that already carry a suffix are folded into the '*' of the same base."""
    bases_all = {parse_query(n)[0].lower() for n in numbers if parse_query(n)[1] == ""}
    out, seen = [], set()
    for n in numbers:
        base, var = parse_query(n)
        if var is None:                       # not an invoice-looking text (e.g. 'korea') -> unchanged
            item = n
        elif base.lower() in bases_all:
            item = base + "-*"
        else:
            item = n
        if item.lower() not in seen:
            seen.add(item.lower())
            out.append(item)
    return out


def cfg_for_variant(cfg, variant):
    """'All sheets' mode: every INV-n / PL-n sheet is printed only if it holds data (blank ones are skipped)."""
    if variant != ALL_VARIANTS:
        return cfg
    c = dict(cfg)
    c["conditional_sheets"] = list(cfg.get("conditional_sheets", [])) + [r"^(INV|PL)\s*-\s*\d+"]
    return c


def no_sheet_message(variant):
    if variant == ALL_VARIANTS:
        return "No INV / PL sheet in this file - nothing to print."
    if variant is not None and not isinstance(variant, str):
        return "No " + " / ".join(f"INV{'-' + v if v else ''}" for v in variant) + " sheet in this file - nothing to print."
    if variant:
        return f"No INV-{variant} / PL-{variant} sheet in this file - nothing to print."
    return "No INV / PL sheet in this file - nothing to print."


def _grid_rows(grid):
    if grid is None:
        return []
    if not isinstance(grid, (list, tuple)):
        grid = [[grid]]
    return [list(r) if isinstance(r, (list, tuple)) else [r] for r in grid]


def sheet_signature(grid):
    """Content fingerprint of a sheet, INCLUDING the invoice number (INV.NO. C264958 / C264958-1 ...).
    A sheet with its own invoice number (e.g. INV-1 = C264958-1) is a different invoice, not a
    duplicate - even when the goods lines look the same."""
    sig = []
    for row in _grid_rows(grid):
        for v in row:
            if v is None:
                continue
            if isinstance(v, str):
                v = re.sub(r"\s+", " ", v).strip().upper()
                if not v:
                    continue
            elif isinstance(v, float):
                v = round(v, 4)
            sig.append(v)
    return tuple(sig)


def _sheet_kind(name):
    n = name.strip().upper()
    return "PL" if n.startswith("PL") else "INV"


def pair_warning(picked):
    """Text when a sheet group has INV but no PL (or PL but no INV), e.g. only INV would come out.
    Groups are by variant: INV + PL, INV-1 + PL-1, ...  Returns "" when every group is a complete INV+PL pair."""
    groups = {}
    for n in picked:
        groups.setdefault(sheet_variant(n), set()).add(_sheet_kind(n))
    bad = []
    for v, kinds in groups.items():
        label = f"-{v}" if v else ""
        if "INV" in kinds and "PL" not in kinds:
            bad.append(f"no PL{label} sheet")
        elif "PL" in kinds and "INV" not in kinds:
            bad.append(f"no INV{label} sheet")
    return ("Only part of the pair will print: " + ", ".join(bad)) if bad else ""


_COPY_SUFFIX_RE = re.compile(r"\s*\(\d+\)\s*$")


def copy_base(name):
    """Group key of a sheet: its name without the Excel copy suffix "(2)", "(3)" ...
    'INV(KNITTED) (2)' -> 'inv(knitted)'   'INV (3)' / 'INV(3)' -> 'inv'
    'INV-1 (2)' -> 'inv-1'  (INV-1 stays a different group from INV, so INV-1 / INV-2 are never merged)."""
    return re.sub(r"\s+", "", _COPY_SUFFIX_RE.sub("", name)).lower()


def drop_numbered_copies(picked, get_grid=None):
    """Excel-renamed copies "INV(KNITTED) (2)", "PL(KNITTED) (2)", "INV (3)" are NOT invoice variants: they are
    skipped whenever the sheet without the number is also being printed (first one in workbook order wins).
    Real variants - INV-1, PL-1, INV-2 ... - have their own name, so they are never touched here.
    Returns (kept, dupes) with dupes = [(name, kept_name)].   (get_grid is unused; kept for compatibility)"""
    first, kept, dupes = {}, [], []
    for n in picked:
        key = copy_base(n)
        if key in first:
            dupes.append((n, first[key]))
        else:
            first[key] = n
            kept.append(n)
    return kept, dupes


def drop_duplicates(picked, cfg, get_grid):
    """picked sheet names -> (kept, dupes).  `dupes` is [(name, duplicate_of)].
    1) "(n)" copies (INV, INV (2), INV (3)) -> only the first is printed; INV-1 / INV-2 are different  [skip_numbered_copies]
    2) identical content (same invoice no. + same lines), INV vs INV, PL vs PL           [skip_duplicate_sheets]
    The first one in workbook order always wins."""
    dupes = []
    if cfg.get("skip_numbered_copies", True):
        picked, dupes = drop_numbered_copies(picked, get_grid)
    if not cfg.get("skip_duplicate_sheets", True):
        return list(picked), dupes
    seen, kept = {}, []
    for n in picked:
        try:
            sig = sheet_signature(get_grid(n))
        except Exception:
            kept.append(n)       # cannot read -> never drop a sheet by mistake
            continue
        key = (_sheet_kind(n), sig)
        if sig and key in seen:
            dupes.append((n, seen[key]))
        else:
            seen[key] = n
            kept.append(n)
    return kept, dupes


def _is_conditional(name, cfg):
    return any(re.search(p, name, re.I) for p in cfg.get("conditional_sheets", []))


def grid_has_data(grid):
    """True when the sheet holds at least one line item.
    `grid` is a 2-D list of cell values. We find the 'QUANTITY' header cell, then look down that
    column for a number > 0 (the TOTAL rows are formulas that evaluate to 0 on an empty sheet, so an
    empty sheet is correctly reported as having no data)."""
    if grid is None:
        return False
    if not isinstance(grid, (list, tuple)):
        grid = [[grid]]
    rows = [list(r) if isinstance(r, (list, tuple)) else [r] for r in grid]
    hdr = None
    for r, row in enumerate(rows[:60]):
        for c, v in enumerate(row):
            if isinstance(v, str) and "QUANTITY" in v.upper():
                hdr = (r, c)
                break
        if hdr:
            break
    if not hdr:
        return True          # unknown layout -> never skip a sheet by mistake
    r0, c0 = hdr
    for row in rows[r0 + 1:]:
        if c0 < len(row):
            v = row[c0]
            if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
                return True
    return False


def split_by_condition(picked, cfg, has_data):
    """picked sheet names -> (to_print, skipped).  `has_data(name)` -> bool is only called for
    conditional sheets (INV-1 / PL-1 / INV-2 / PL-2)."""
    to_print, skipped = [], []
    for n in picked:
        if _is_conditional(n, cfg) and not has_data(n):
            skipped.append(n)
        else:
            to_print.append(n)
    return to_print, skipped


def plan_sheets_from_file(path, cfg, variant=None):
    """Fast look at sheet names without starting Excel (xlsx/xlsm only).
    Returns (all names, sheets that will print, empty conditional sheets skipped,
    duplicate sheets skipped as [(name, duplicate_of)])."""
    if Path(path).suffix.lower() not in (".xlsx", ".xlsm"):
        return None
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        picked = filter_by_variant(pick_sheets(names, cfg["sheet_include"], cfg["sheet_exclude"]), variant)
        cache = {}

        def get_grid(name):
            if name not in cache:
                cache[name] = [list(r) for r in wb[name].iter_rows(max_row=80, values_only=True)]
            return cache[name]

        ccfg = cfg_for_variant(cfg, variant)
        to_print, empty = split_by_condition(picked, ccfg, lambda n: grid_has_data(get_grid(n)))
        to_print, dupes = drop_duplicates(to_print, cfg, get_grid)
    finally:
        wb.close()
    return names, to_print, empty, dupes


def prepare_copy(path, tmp_dir, idx, fixes):
    """Copy of the workbook with mis-spelled fonts corrected (xlsx/xlsm only; lossless zip patch).
    Returns the path Excel should open (the original if nothing needed fixing)."""
    import zipfile
    if not fixes or Path(path).suffix.lower() not in (".xlsx", ".xlsm"):
        return path
    pairs = [(k.encode("utf-8"), v.encode("utf-8")) for k, v in fixes.items() if k and k != v]
    try:
        with zipfile.ZipFile(path) as zin:
            hit = False
            payload = []
            for it in zin.infolist():
                data = zin.read(it.filename)
                if it.filename.lower().endswith(".xml") and it.filename.startswith("xl/"):
                    new = data
                    for a, b in pairs:
                        new = new.replace(b'"' + a + b'"', b'"' + b + b'"')
                    hit = hit or new != data
                    data = new
                payload.append((it, data))
            if not hit:
                return path
            out = os.path.join(tmp_dir, f"{idx:04d}{Path(path).suffix.lower()}")
            with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
                for it, data in payload:
                    zout.writestr(it, data)
        return out
    except Exception:
        return path   # never block the job because of the font patch


def pdf_layout_printer(chosen, cfg):
    """Printer whose metrics Excel uses to lay out a PDF export (see config pdf_layout_printer)."""
    want = (cfg.get("pdf_layout_printer") or chosen or "").strip()
    if want:
        return want
    try:
        import win32print
        win32print.GetDefaultPrinter()
        return ""                                   # a real default printer exists -> keep it
    except Exception:
        pass
    for p in list_printers():
        if "microsoft print to pdf" in p.lower():
            return p
    return ""


def find_pdf_printer():
    for p in list_printers():
        if "microsoft print to pdf" in p.lower():
            return p
    return ""


def _wait_for_file(path, timeout=90):
    """Wait until the PDF printer has finished writing `path` (exists, > 0 bytes, size stable)."""
    import time
    end = time.time() + timeout
    last = -1
    while time.time() < end:
        if os.path.exists(path):
            size = os.path.getsize(path)
            if size > 0 and size == last:
                return True
            last = size
        time.sleep(0.5)
    return False


def make_pdf(excel, out_pdf, engine):
    """Selected sheets of the active workbook -> out_pdf. Returns the engine that really produced it."""
    if os.path.exists(out_pdf):
        try:
            os.remove(out_pdf)
        except OSError:
            pass
    if engine == "printer":
        try:
            excel.ActiveSheet.PrintOut(Copies=1, PrintToFile=True, PrToFileName=str(out_pdf))
            if _wait_for_file(str(out_pdf)):
                return "printer"
        except Exception:
            pass
        # fall through -> Excel's own exporter
    excel.ActiveSheet.ExportAsFixedFormat(0, str(out_pdf), 0, True, False)
    return "export"


# --------------------------------------------------------------------------
# Make EXCEL itself print on the chosen printer
# (changing the Windows default is not enough: Excel keeps its own "active printer")
# --------------------------------------------------------------------------
def _printer_ports(name):
    try:
        import win32print
        h = win32print.OpenPrinter(name)
        try:
            raw = str(win32print.GetPrinter(h, 2).get("pPortName", ""))
        finally:
            win32print.ClosePrinter(h)
        return [x.strip().rstrip(":") for x in raw.split(",") if x.strip()]
    except Exception:
        return []


def _active_matches(active, name):
    a, n = (active or "").strip().lower(), (name or "").strip().lower()
    return bool(n) and (a == n or a.startswith(n + " on "))


def select_excel_printer(excel, name, ports=None):
    """Set excel.ActivePrinter to `name` and VERIFY it. Returns the ActivePrinter text Excel reports.
    Excel wants 'Printer name on Ne03:' - the 'NeXX' part differs per PC, so we try the real port first,
    then Ne00..Ne99.  Raises RuntimeError when Excel cannot be switched."""
    try:
        cur = excel.ActivePrinter
    except Exception:
        cur = ""
    if _active_matches(cur, name):
        return cur
    if ports is None:
        ports = _printer_ports(name)
    candidates = [f"{name} on {p}:" for p in ports] + [name] + [f"{name} on Ne{n:02d}:" for n in range(100)]
    for cand in candidates:
        try:
            excel.ActivePrinter = cand
        except Exception:
            continue
        try:
            now = excel.ActivePrinter
        except Exception:
            now = cand
        if _active_matches(now, name):
            return now
    raise RuntimeError(f"Excel would not switch to the printer '{name}' (Excel is currently on '{cur or 'unknown'}'). "
                       "Nothing was printed, so it cannot come out on the wrong printer. "
                       "Check the printer is installed/online for this Windows user, or set "
                       "\"strict_printer\": false in the config to print anyway.")


# --------------------------------------------------------------------------
# Excel work (print / pdf) - handles one OR many files in a single Excel session
# --------------------------------------------------------------------------
def process_files(paths, cfg, mode, copies=1, printer="", pdf_path=None, source="manual", progress=None,
                  variants=None, pdf_paths=None):
    """mode = 'print' | 'pdf'.  pdf_paths = one output PDF per file (used by the combine feature).  variants = [None | '' | '1', ...] aligned with paths -> see filter_by_variant.  Never raises for a single bad file: returns one dict per file
    {path, name, ok, sheets, error}.  Every file is written to the print log."""
    import pythoncom
    import shutil
    import tempfile
    import win32com.client

    tag = f"{source}/{mode}"
    results = []

    def log(path, ok, sheets, msg=""):
        name = os.path.basename(path)
        write_log(tag, name, sheets, copies if mode == "print" else "", printer, "OK" if ok else "FAILED", msg)
        results.append({"path": path, "name": name, "ok": ok, "sheets": sheets, "error": msg})

    with JOB_LOCK:
        pythoncom.CoInitialize()
        excel = None
        restore_printer = None
        try:
            try:
                engine = cfg.get("pdf_engine", "printer") if mode == "pdf" else ""
                pdf_prt = find_pdf_printer() if engine == "printer" else ""
                if engine == "printer" and not pdf_prt:
                    engine = "export"
                if mode == "print":
                    layout_printer = printer
                elif engine == "printer":
                    layout_printer = pdf_prt
                else:
                    layout_printer = pdf_layout_printer(printer, cfg)
                if layout_printer:
                    restore_printer = _switch_default_printer(layout_printer)
                excel = win32com.client.DispatchEx("Excel.Application")
                excel.Visible = False
                excel.DisplayAlerts = False
                for attr in ("ScreenUpdating", "EnableEvents"):   # speed: no redraw / macros while exporting
                    try:
                        setattr(excel, attr, False)
                    except Exception:
                        pass
            except Exception as exc:  # noqa: BLE001
                for p in paths:
                    log(p, False, [], f"Could not start Excel / select printer: {exc}")
                return results

            fix_dir = tempfile.mkdtemp(prefix="invfix_")
            total = len(paths)
            for i, path in enumerate(paths, 1):
                wb = None
                picked = []
                name = os.path.basename(path)
                try:
                    if progress:
                        progress(f"{'Printing' if mode == 'print' else 'Exporting'} {name}  ({i}/{total})…")
                    open_path = prepare_copy(path, fix_dir, i, cfg.get("font_fixes")) if mode == "pdf" else path
                    wb = excel.Workbooks.Open(str(open_path), UpdateLinks=0, ReadOnly=True)
                    names = [ws.Name for ws in wb.Worksheets if ws.Visible == -1]  # -1 = visible
                    picked = pick_sheets(names, cfg["sheet_include"], cfg["sheet_exclude"])
                    picked = filter_by_variant(picked, (variants[i - 1] if variants else None))
                    if not picked:
                        v = variants[i - 1] if variants else None
                        raise RuntimeError((no_sheet_message(v) if v is not None else
                                            "No sheet matched the INV/PL rules.")
                                           + " Sheets in file: " + ", ".join(names))

                    def has_data(sheet_name):
                        try:
                            return grid_has_data(wb.Worksheets(sheet_name).UsedRange.Value)
                        except Exception:
                            return True  # cannot read -> print rather than silently drop

                    picked, skipped_empty = split_by_condition(
                        picked, cfg_for_variant(cfg, variants[i - 1] if variants else None), has_data)
                    picked, dupes = drop_duplicates(
                        picked, cfg, lambda n: wb.Worksheets(n).UsedRange.Value)
                    notes = []
                    if skipped_empty:
                        notes.append("Skipped empty: " + ", ".join(skipped_empty))
                    if dupes:
                        notes.append("Skipped duplicate: " + ", ".join(f"{a} (= {b})" for a, b in dupes))
                    warn = pair_warning(picked)
                    if warn:
                        notes.append("WARNING: " + warn)
                    note = "; ".join(notes)
                    if not picked:
                        raise RuntimeError("Nothing to print - all matching sheets are empty.")

                    if mode == "print":
                        used_printer = ""
                        if printer:
                            try:
                                used_printer = select_excel_printer(excel, printer)
                            except RuntimeError as exc:
                                if cfg.get("strict_printer", True):
                                    raise
                                note = (note + "; " if note else "") + "WARNING: " + str(exc)
                        if not used_printer:
                            try:
                                used_printer = excel.ActivePrinter
                            except Exception:
                                used_printer = "(unknown)"
                        note = (note + "; " if note else "") + f"Excel printer: {used_printer}"
                        grouped = False
                        if cfg.get("print_one_job_per_file", True) and len(picked) > 1:
                            try:   # select INV + PL together -> a single print job ...
                                wb.Worksheets(picked[0]).Select()
                                for sheet in picked[1:]:
                                    wb.Worksheets(sheet).Select(False)
                                # ... but ONLY if Excel really has all of them selected (otherwise just INV would print)
                                grouped = int(excel.ActiveWindow.SelectedSheets.Count) == len(picked)
                            except Exception:
                                grouped = False
                        if grouped:
                            excel.ActiveWindow.SelectedSheets.PrintOut(Copies=int(copies), Collate=True)
                            note = (note + "; " if note else "") + "1 print job"
                        else:
                            for sheet in picked:  # workbook order: INV first, then PL
                                wb.Worksheets(sheet).PrintOut(Copies=int(copies), Collate=True)
                        note = (note + "; " if note else "") + "Printed: " + ", ".join(picked)
                    else:
                        wb.Worksheets(picked[0]).Select()
                        for sheet in picked[1:]:
                            wb.Worksheets(sheet).Select(False)  # False = add to selection
                        out_pdf = pdf_paths[i - 1] if pdf_paths else pdf_path
                        used = make_pdf(excel, out_pdf, engine)
                        if used != engine:
                            note = (note + "; " if note else "") + "PDF engine fell back to Excel export"
                    log(path, True, picked, note)
                except Exception as exc:  # noqa: BLE001
                    log(path, False, picked, str(exc))
                finally:
                    if wb is not None:
                        try:
                            wb.Close(SaveChanges=False)
                        except Exception:
                            pass
            return results
        finally:
            shutil.rmtree(locals().get("fix_dir") or "", ignore_errors=True)
            if excel is not None:
                try:
                    excel.Quit()
                except Exception:
                    pass
            if restore_printer:
                try:
                    import win32print
                    win32print.SetDefaultPrinter(restore_printer)
                except Exception:
                    pass
            pythoncom.CoUninitialize()


# --------------------------------------------------------------------------
# Print existing PDF files (PDF mode)
# --------------------------------------------------------------------------
def _first_existing(paths):
    for p_ in paths:
        if p_ and os.path.isfile(p_):
            return p_
    return ""


def find_pdf_viewer(cfg):
    """Program used to print PDFs silently. Returns (kind, exe):
         ("sumatra", exe)  best: prints to a NAMED printer, quiet, supports copies
         ("adobe", exe)    Acrobat / Reader  (/t file printer)
         ("shell", "")     Windows 'printto' verb of whatever PDF program is the default (works with Adobe/Foxit, not Edge)
    Put SumatraPDF.exe next to InvoicePrinter.exe (or in a 'tools' folder) and it is picked up automatically."""
    custom = (cfg.get("pdf_viewer_path") or "").strip().strip('"')
    if custom and os.path.isfile(custom):
        low = os.path.basename(custom).lower()
        if "sumatra" in low:
            return ("sumatra", custom)
        if "acro" in low:
            return ("adobe", custom)
    pf = [os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")]
    local = os.environ.get("LOCALAPPDATA", "")
    import shutil
    import glob
    # The current portable download is named e.g. "SumatraPDF-3.5.2-64.exe" / "SumatraPDF-x.y-portable.exe",
    # not "SumatraPDF.exe" - so accept any SumatraPDF*.exe in the app folder, its 'tools' folder or Downloads.
    dl = os.path.join(os.path.expanduser("~"), "Downloads")
    found_named = []
    for folder in (str(APP_DIR), str(APP_DIR / "tools"), dl):
        found_named += sorted(glob.glob(os.path.join(folder, "SumatraPDF*.exe")), reverse=True)
    found_named = [f for f in found_named if "install" not in os.path.basename(f).lower()]
    sumatra = _first_existing([
        # a really installed SumatraPDF first (a loose SumatraPDF*.exe may be the installer renamed)
        os.path.join(local, "SumatraPDF", "SumatraPDF.exe") if local else "",
        os.path.join(pf[0], "SumatraPDF", "SumatraPDF.exe"), os.path.join(pf[1], "SumatraPDF", "SumatraPDF.exe"),
        str(APP_DIR / "SumatraPDF.exe"), str(APP_DIR / "tools" / "SumatraPDF.exe"), *found_named,
        shutil.which("SumatraPDF") or ""])
    if sumatra:
        return ("sumatra", sumatra)
    adobe = _first_existing(
        [os.path.join(b, *rest) for b in pf for rest in (
            ("Adobe", "Acrobat DC", "Acrobat", "Acrobat.exe"),
            ("Adobe", "Acrobat Reader DC", "Reader", "AcroRd32.exe"),
            ("Adobe", "Acrobat Reader", "Reader", "AcroRd32.exe"),
            ("Adobe", "Reader 11.0", "Reader", "AcroRd32.exe"))])
    if adobe:
        return ("adobe", adobe)
    return ("shell", "")


def build_pdf_print_command(kind, exe, path, printer, copies, cfg=None):
    """Command line (list) for one PDF - or None for kind 'shell'."""
    path = os.path.normpath(path)
    if kind == "sumatra":
        cmd = [exe, "-print-to", printer] if printer else [exe, "-print-to-default"]
        cfg = cfg or {}
        settings = ",".join(x for x in ((cfg.get("pdf_sumatra_extra") or "").strip(), f"{max(1, int(copies))}x") if x)
        cmd += ["-print-settings", settings]
        if cfg.get("pdf_sumatra_silent", False):
            cmd.append("-silent")
        return cmd + ["-exit-when-done", path]
    if kind == "adobe":
        return [exe, "/t", path] + ([printer] if printer else [])
    return None


def _pdf_has_printto_verb():
    """True only if the default PDF program really registers a 'printto' verb (Edge / Chrome do NOT)."""
    try:
        import winreg
        prog = winreg.QueryValue(winreg.HKEY_CLASSES_ROOT, ".pdf")
        winreg.QueryValue(winreg.HKEY_CLASSES_ROOT, prog + r"\shell\printto\command")
        return True
    except Exception:
        return False


def _printer_exists(name):
    """False when the chosen printer name is not installed on this PC (so we can say so instead of failing silently)."""
    if not name:
        return True
    try:
        import win32print
        flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
        return any(p[2].lower() == name.lower() for p in win32print.EnumPrinters(flags))
    except Exception:
        return True          # cannot check -> do not block


def _spool_jobs(printer):
    """Jobs currently in the Windows print queue of this printer (None if it cannot be read)."""
    try:
        import win32print
        h = win32print.OpenPrinter(printer or win32print.GetDefaultPrinter())
        try:
            return win32print.EnumJobs(h, 0, -1, 1)
        finally:
            win32print.ClosePrinter(h)
    except Exception:
        return None


def _spool_report(printer, path, seconds=10):
    """After the PDF program returns, watch the Windows queue for this file's job.
    Raises RuntimeError when the job is stuck / in error (so the log no longer says OK), else returns a note."""
    import time
    try:
        import win32print as wp
    except Exception:
        return ""
    stem = Path(path).stem.lower()
    bad = [("JOB_STATUS_ERROR", "error"), ("JOB_STATUS_OFFLINE", "printer offline"),
           ("JOB_STATUS_PAPEROUT", "paper out"), ("JOB_STATUS_USER_INTERVENTION", "printer needs attention"),
           ("JOB_STATUS_BLOCKED_DEVQ", "blocked (printer/driver problem)"), ("JOB_STATUS_PAUSED", "queue paused")]
    seen = False
    end = time.time() + seconds
    while time.time() < end:
        jobs = _spool_jobs(printer)
        if jobs is None:
            return " | queue not readable"
        mine = [j for j in jobs if stem in str(j.get("pDocument", "")).lower()]
        for j in mine:
            seen = True
            for const, text in bad:
                if j.get("Status", 0) & getattr(wp, const, 0):
                    raise RuntimeError(f"Job reached the Windows print queue but is stuck: {text}. "
                                       f"Open Settings > Printers > this printer > Open queue and check it.")
        if seen and not mine:
            return " | job left the queue (sent to printer)"
        if seen:
            time.sleep(0.5)
            continue
        time.sleep(0.5)
    if seen:
        return " | job still in queue after %ds (printer busy / waiting)" % seconds
    return " | no job seen in Windows queue (printed instantly, or it never reached the spooler)"


def _print_one_pdf(kind, exe, path, printer, copies, cfg=None):
    """Send ONE pdf to the printer. Returns a short text for the log, raises RuntimeError on failure."""
    import subprocess
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    if not _printer_exists(printer):
        raise RuntimeError(f"Printer not found on this PC: '{printer}'. Pick another printer in the drop-down.")
    cmd = build_pdf_print_command(kind, exe, path, printer, copies, cfg)
    if kind == "sumatra":
        try:
            r = subprocess.run(cmd, creationflags=flags, timeout=60, capture_output=True, text=True)
        except subprocess.TimeoutExpired:
            try:
                subprocess.run(["taskkill", "/F", "/T", "/IM", os.path.basename(exe)], creationflags=flags,
                               capture_output=True)
            except Exception:
                pass
            raise RuntimeError(
                f"SumatraPDF did not finish within 60 s, so nothing was printed. If an 'Installer' window opened, "
                f"'{exe}' is the SumatraPDF INSTALLER, not the program: click Install once, then restart this app "
                f"(or set \"pdf_viewer_path\" to %LOCALAPPDATA%\\SumatraPDF\\SumatraPDF.exe).")
        if r.returncode != 0:
            raise RuntimeError(f"SumatraPDF could not print (code {r.returncode}) {r.stderr.strip()[:200]}")
        note = _spool_report(printer, path)
        return f"SumatraPDF ({os.path.basename(exe)}) -> {printer or '(default)'}{note}"
    if kind == "adobe":
        for _ in range(max(1, int(copies))):   # Adobe has no copies switch
            proc = subprocess.Popen(cmd, creationflags=flags)
            try:
                proc.wait(timeout=25)
            except subprocess.TimeoutExpired:
                pass                           # Reader stays open after printing - the job is already sent
        return f"Adobe ({os.path.basename(exe)}) -> {printer or '(default)'}"
    if not _pdf_has_printto_verb():
        raise RuntimeError(
            "Nothing was sent to the printer: no PDF printing program was found (SumatraPDF / Adobe Reader), and the "
            "default PDF app (Edge/Chrome) cannot print silently. Fix: download the portable SumatraPDF.exe and put it "
            "next to this program (or in a 'tools' folder), or set \"pdf_viewer_path\" in invoice_printer_config.json.")
    try:                                       # generic: the default PDF program's "printto" verb
        import win32api
        for _ in range(max(1, int(copies))):
            win32api.ShellExecute(0, "printto", os.path.normpath(path), f'"{printer}"' if printer else None, ".", 0)
        return f"Windows printto -> {printer or '(default)'}"
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(
            "No program found that can print PDFs silently (" + str(exc) + "). Fix: put SumatraPDF.exe "
            "(free, portable) next to InvoicePrinter.exe, or install Adobe Reader, or set \"pdf_viewer_path\" in the config.")


def print_pdf_files(paths, cfg, copies=1, printer="", source="manual", progress=None):
    """Print PDF files exactly as they are. Same result shape as process_files; every file goes to the print log."""
    tag = f"{source}/print-pdf"
    results = []

    def log(path, ok, msg=""):
        name = os.path.basename(path)
        write_log(tag, name, ["PDF file"], copies, printer, "OK" if ok else "FAILED", msg)
        results.append({"path": path, "name": name, "ok": ok, "sheets": ["PDF file"], "error": msg})

    with JOB_LOCK:
        kind, exe = find_pdf_viewer(cfg)
        total = len(paths)
        for i, path in enumerate(paths, 1):
            try:
                if progress:
                    progress(f"Printing {os.path.basename(path)}  ({i}/{total})…")
                if not os.path.isfile(path):
                    raise RuntimeError("File not found: " + path)
                log(path, True, _print_one_pdf(kind, exe, path, printer, copies, cfg))
            except Exception as exc:  # noqa: BLE001
                log(path, False, str(exc))
    return results


# --------------------------------------------------------------------------
# Combine many PDFs into ONE (so you print a single file instead of N jobs)
# --------------------------------------------------------------------------
def merge_pdfs(pdf_files, out_path, titles=None):
    """Merge pdf_files (in order) into out_path. Adds one bookmark per source file.
    Needs:  pip install pypdf"""
    try:
        from pypdf import PdfWriter, PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfWriter, PdfReader  # older name
        except ImportError:
            raise RuntimeError("Combining PDFs needs the 'pypdf' package, which is missing from this build.\n"
                               "Developers: run   pip install pypdf   and rebuild.")
    writer = PdfWriter()
    for n, f in enumerate(pdf_files):
        reader = PdfReader(str(f))
        first = len(writer.pages)
        for page in reader.pages:
            writer.add_page(page)
        if len(reader.pages):
            try:
                writer.add_outline_item((titles[n] if titles else Path(f).stem), first)
            except Exception:
                pass
    with open(out_path, "wb") as fh:
        writer.write(fh)
    return len(writer.pages)


def export_and_combine(rows, cfg, out_path, progress=None):
    """rows = [{path, name, variant}, ...] -> ONE pdf at out_path. Returns (pages, ok_count, failed_list)."""
    import shutil
    import tempfile
    tmp = tempfile.mkdtemp(prefix="invprint_")
    try:
        paths = [r["path"] for r in rows]
        variants = [r.get("variant") for r in rows]
        outs = [os.path.join(tmp, f"{i:04d}.pdf") for i in range(len(paths))]
        res = process_files(paths, cfg, "pdf", source="combine", progress=progress,
                            variants=variants, pdf_paths=outs)
        good = [(rows[i], outs[i]) for i, r in enumerate(res) if r["ok"] and os.path.exists(outs[i])]
        bad = [r for r in res if not r["ok"]]
        for i, r in enumerate(res):   # Excel said OK but no PDF appeared -> report it instead of hiding it
            if r["ok"] and not os.path.exists(outs[i]):
                bad.append({**r, "error": "PDF file was not produced by Excel / the PDF printer."})
        if not good:
            raise RuntimeError("No PDF could be created:\n" + "\n".join(f"• {r['name']}: {r['error']}" for r in bad))
        if progress:
            progress(f"Merging {len(good)} PDF(s)…")
        pages = merge_pdfs([o for _, o in good], out_path,
                           [Path(r["name"]).stem for r, _ in good])
        return pages, len(good), bad
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# --------------------------------------------------------------------------
# Auto-print watcher (background thread, polling - no extra libraries needed)
# --------------------------------------------------------------------------
def _sig(path):
    try:
        st = os.stat(path)
        return (st.st_mtime, st.st_size)
    except OSError:
        return None


def _sha1(path):
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def scan_excel(folder, cfg):
    """All Excel files under `folder` that pass the name filter."""
    exts = tuple(e.lower() for e in cfg["extensions"])
    rx = re.compile(cfg["watch_name_regex"], re.I) if cfg.get("watch_name_regex") else None
    found = []
    for dirpath, _dirs, files in os.walk(folder):
        for name in files:
            if name.startswith("~$") or not name.lower().endswith(exts):
                continue
            if rx and not rx.search(name):
                continue
            found.append(os.path.join(dirpath, name))
    return found


def watch_loop(stop, cfg, emit):
    """Print every NEW Excel file that appears in cfg['watch_folder'].
    * Files that already exist when watching starts are ignored (baseline).
    * A file is printed only after its size/time stayed unchanged for one full check
      (so half-downloaded Telegram files are never printed).
    * The same file content is never printed twice (SHA-1 kept in auto_print_state.json)."""
    folder = cfg.get("watch_folder") or cfg.get("root_folder")
    if not folder or not os.path.isdir(folder):
        emit("⚠ Watch folder not found: " + str(folder))
        return

    seen = {}      # path -> signature already handled (or baseline)
    pending = {}   # path -> signature seen once, waiting to be stable
    try:
        for p in scan_excel(folder, cfg):
            seen[p] = _sig(p)
    except Exception as exc:  # noqa: BLE001
        emit(f"⚠ Cannot read watch folder: {exc}")
        return
    state = load_state()
    emit(f"👁 Auto-print ON – watching {folder}")

    while not stop.wait(max(2, int(cfg.get("watch_interval_sec", 5)))):
        try:
            files = scan_excel(folder, cfg)
        except Exception:
            continue
        for p in files:
            sig = _sig(p)
            if sig is None or seen.get(p) == sig:
                continue
            if pending.get(p) != sig:          # new or still changing -> check again next round
                pending[p] = sig
                continue
            pending.pop(p, None)
            seen[p] = sig
            name = os.path.basename(p)
            try:
                digest = _sha1(p)
            except OSError:
                seen.pop(p, None)              # locked/unreadable -> retry later
                continue
            if digest in state["hashes"]:
                write_log("auto/print", name, [], "", cfg.get("printer"), "SKIPPED", "Same file already printed")
                emit(f"⏭ Skipped duplicate: {name}")
                continue
            emit(f"🖨 Auto-printing {name}…")
            try:
                res = process_files([p], cfg, "print", copies=int(cfg.get("copies", 1)),
                                    printer=cfg.get("printer", ""), source="auto")[0]
            except Exception as exc:  # noqa: BLE001  (e.g. pywin32 missing)
                write_log("auto/print", name, [], "", cfg.get("printer"), "FAILED", str(exc))
                emit(f"❌ Auto-print failed: {exc}")
                continue
            if res["ok"]:
                state["hashes"].append(digest)
                state["hashes"] = state["hashes"][-2000:]
                save_state(state)
                emit(f"✅ Auto-printed {name}: " + ", ".join(res["sheets"]))
            else:
                emit(f"❌ {name}: {res['error']}")
    emit("Auto-print OFF")


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------
class App:
    def __init__(self, root):
        self.root = root
        self.cfg = load_config()
        if not self.cfg["watch_folder"]:
            self.cfg["watch_folder"] = self.cfg["root_folder"]
        self.results = []
        self.q = queue.Queue()
        self.cancel = threading.Event()
        self.busy = False
        self.plan = {}  # path -> (all names, picked) or None
        self.watch_stop = None
        self.batch_win = None
        self.batch_rows = []

        root.title(f"{APP_NAME}  v{APP_VERSION}")
        root.geometry("1000x760")
        root.minsize(860, 600)
        self._style()
        self._build()
        root.after(100, self._poll)
        root.after(300, self._first_run)
        if self.cfg.get("auto_print_enabled") and self.cfg["watch_folder"]:
            self.auto_var.set(True)
            self.start_watch()
        root.protocol("WM_DELETE_WINDOW", self.on_close)

    def _first_run(self):
        """Welcome + setup on a fresh machine (or when the saved folder doesn't exist here)."""
        if not excel_installed():
            messagebox.showwarning(
                "Microsoft Excel not found",
                "Invoice Printer prints through desktop Microsoft Excel, but Excel was not found on this PC.\n\n"
                "You can still open the app, but printing will fail until Excel is installed.")
        root_dir = self.cfg.get("root_folder", "")
        if root_dir and os.path.isdir(root_dir):
            return
        first = not root_dir
        messagebox.showinfo(
            "Welcome" if first else "Folder not found",
            ("Welcome to Invoice Printer!\n\n" if first else
             f"The saved invoice folder does not exist on this PC:\n{root_dir}\n\n")
            + "Please choose the top folder that contains your month folders.\n"
              "You can change it later in the main window or in Settings.")
        self.browse_root()
        wf = self.cfg.get("watch_folder", "")
        if wf and not os.path.isdir(wf):          # stale watch folder from another PC
            self.cfg["watch_folder"] = self.cfg.get("root_folder", "")
            self.cfg["auto_print_enabled"] = False
            save_config(self.cfg)
            try:
                if self.watch_stop:
                    self.watch_stop.set()
                    self.watch_stop = None
                self.auto_var.set(False)
                self.watch_var.set(self.cfg["watch_folder"])
            except Exception:
                pass

    def on_close(self):
        if self.watch_stop:
            self.watch_stop.set()
        self.root.destroy()

    # ---- look & feel
    def _style(self):
        style = ttk.Style()
        for theme in ("vista", "clam"):
            if theme in style.theme_names():
                style.theme_use(theme)
                break
        base = ("Segoe UI", 10)
        style.configure(".", font=base)
        style.configure("Title.TLabel", font=("Segoe UI Semibold", 16))
        style.configure("Hint.TLabel", foreground="#666666")
        style.configure("Plan.TLabel", foreground="#0b6b2e", font=("Segoe UI Semibold", 10))
        style.configure("Primary.TButton", font=("Segoe UI Semibold", 10), padding=(14, 6))
        style.configure("Treeview", rowheight=26)
        style.configure("Treeview.Heading", font=("Segoe UI Semibold", 10))

    def _build(self):
        pad = {"padx": 16}
        ttk.Label(self.root, text="Invoice Printer", style="Title.TLabel").pack(anchor="w", pady=(14, 0), **pad)
        ttk.Label(self.root, text="Search an invoice number, then print only the INV and PL sheets.",
                  style="Hint.TLabel").pack(anchor="w", pady=(0, 10), **pad)

        # --- folder + printer
        top = ttk.Frame(self.root)
        top.pack(fill="x", **pad)
        ttk.Label(top, text="Root folder").grid(row=0, column=0, sticky="w")
        self.root_var = tk.StringVar(value=self.cfg["root_folder"])
        ttk.Entry(top, textvariable=self.root_var, state="readonly").grid(row=0, column=1, sticky="ew", padx=8)
        ttk.Button(top, text="Browse…", command=self.browse_root).grid(row=0, column=2)
        ttk.Button(top, text="Settings", command=self.open_settings).grid(row=0, column=3, padx=(6, 0))

        ttk.Label(top, text="Printer").grid(row=1, column=0, sticky="w", pady=(8, 0))
        printers = [DEFAULT_PRINTER_LABEL] + list_printers()
        self.printer_var = tk.StringVar(value=self.cfg["printer"] if self.cfg["printer"] in printers
                                        else DEFAULT_PRINTER_LABEL)
        self.printer_box = ttk.Combobox(top, textvariable=self.printer_var, values=printers, state="readonly")
        self.printer_box.grid(row=1, column=1, sticky="ew", padx=8, pady=(8, 0))
        self.printer_box.bind("<<ComboboxSelected>>", lambda _e: self.on_printer())
        ttk.Button(top, text="Refresh", command=self.refresh_printers).grid(row=1, column=2, pady=(8, 0))
        ttk.Button(top, text="Printer setup…", command=self.open_printer_setup).grid(
            row=1, column=3, padx=(6, 0), pady=(8, 0))
        top.columnconfigure(1, weight=1)

        # --- search bar
        bar = ttk.Frame(self.root)
        bar.pack(fill="x", pady=(12, 8), **pad)
        ttk.Label(bar, text="Invoice no.").pack(side="left")
        # multi-line box: paste a whole list (one per line, or separated by , ; space) - every line stays visible
        qbox = ttk.Frame(bar)
        qbox.pack(side="left", fill="x", expand=True, padx=8)
        self.entry = tk.Text(qbox, height=3, wrap="word", font=("Segoe UI", 11), undo=True,
                             relief="solid", borderwidth=1, padx=6, pady=4)
        qscroll = ttk.Scrollbar(qbox, orient="vertical", command=self.entry.yview)
        self.entry.configure(yscrollcommand=qscroll.set)
        self.entry.pack(side="left", fill="x", expand=True)
        qscroll.pack(side="left", fill="y")
        self.entry.bind("<Return>", lambda _e: (self.search(), "break")[1])        # Enter = search
        self.entry.bind("<Shift-Return>", lambda _e: None)                          # Shift+Enter = new line
        self.entry.bind("<KeyRelease>", lambda _e: self.update_count())
        self.entry.bind("<<Paste>>", lambda _e: self.root.after(50, self.update_count))
        self.entry.focus_set()
        side = ttk.Frame(bar)
        side.pack(side="left")
        ttk.Button(side, text="Search", style="Primary.TButton", command=self.search).grid(row=0, column=0)
        ttk.Button(side, text="Batch…", command=self.open_batch).grid(row=0, column=1, padx=(8, 0))
        ttk.Button(side, text="Import Excel…", command=self.import_excel).grid(row=1, column=0,
                                                                               sticky="ew", pady=(6, 0))
        ttk.Button(side, text="Clear", command=self.clear_query).grid(row=1, column=1, sticky="ew",
                                                                      padx=(8, 0), pady=(6, 0))
        opts = ttk.Frame(self.root)
        opts.pack(fill="x", padx=16)
        self.all_var = tk.BooleanVar(value=bool(self.cfg.get("all_variants", False)))
        ttk.Checkbutton(opts, text="All sheets – print INV, INV-1, INV-2 … INV-n (blank ones skipped), no need to type -1, -2",
                        variable=self.all_var, command=self.on_all_toggle).pack(side="left")
        self.all_chk = opts.winfo_children()[-1]   # the "All sheets" checkbox (disabled in PDF mode)
        self.pdf_var = tk.BooleanVar(value=bool(self.cfg.get("print_pdf_files", False)))
        ttk.Checkbutton(opts, text="Print PDF files (search .pdf) instead of Excel",
                        variable=self.pdf_var, command=self.on_pdf_toggle).pack(side="left", padx=(20, 0))
        self.count_var = tk.StringVar(value="")
        ttk.Label(opts, textvariable=self.count_var, style="Hint.TLabel").pack(side="right")

        # --- bottom area (packed bottom-up)
        self.status = tk.StringVar(value="Ready")
        ttk.Label(self.root, textvariable=self.status, style="Hint.TLabel", relief="sunken", anchor="w",
                  padding=(10, 4), wraplength=940).pack(fill="x", side="bottom")

        # auto-print panel
        wf = ttk.LabelFrame(self.root, text=" Auto-print (watch folder) ", padding=(10, 6))
        wf.pack(fill="x", side="bottom", pady=(0, 8), **pad)
        self.watch_var = tk.StringVar(value=self.cfg["watch_folder"])
        ttk.Entry(wf, textvariable=self.watch_var, state="readonly").grid(row=0, column=0, sticky="ew")
        ttk.Button(wf, text="Browse…", command=self.browse_watch).grid(row=0, column=1, padx=(8, 0))
        self.auto_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(wf, text="Print new files automatically", variable=self.auto_var,
                        command=self.toggle_auto).grid(row=0, column=2, padx=(14, 0))
        self.watch_status = tk.StringVar(value="Off")
        ttk.Label(wf, textvariable=self.watch_status, style="Hint.TLabel").grid(
            row=1, column=0, columnspan=3, sticky="w", pady=(4, 0))
        wf.columnconfigure(0, weight=1)

        actions = ttk.Frame(self.root)
        actions.pack(fill="x", side="bottom", pady=(8, 8), **pad)
        ttk.Label(actions, text="Copies").pack(side="left")
        self.copies_var = tk.IntVar(value=int(self.cfg["copies"]))
        ttk.Spinbox(actions, from_=1, to=20, width=4, textvariable=self.copies_var).pack(side="left", padx=(6, 16))
        self.btn_print = ttk.Button(actions, text="Print INV + PL", style="Primary.TButton", command=self.do_print)
        self.btn_print.configure(text=self.print_btn_text())          # PDF mode remembered from last time
        if self.pdf_mode():
            self.all_chk.configure(state="disabled")
        self.btn_print.pack(side="left")
        self.btn_print_all = ttk.Button(actions, text="Print ALL found", command=self.do_print_many,
                                        state="disabled")
        self.btn_print_all.pack(side="left", padx=(8, 0))
        ttk.Button(actions, text="Select all", command=self.select_all).pack(side="left", padx=(8, 0))
        self.btn_pdf = ttk.Button(actions, text="Save as PDF", command=self.do_pdf)
        self.btn_pdf.pack(side="left", padx=8)
        self.btn_combine = ttk.Button(actions, text="Combine → 1 PDF", command=self.do_combine_pdf)
        self.btn_combine.pack(side="left", padx=(0, 8))
        self.btn_open = ttk.Button(actions, text="Open file", command=self.open_file)
        self.btn_open.pack(side="left")
        ttk.Button(actions, text="Print log", command=self.open_log).pack(side="right")

        self.plan_var = tk.StringVar(value="")
        ttk.Label(self.root, textvariable=self.plan_var, style="Plan.TLabel", wraplength=900).pack(
            anchor="w", side="bottom", pady=(8, 0), **pad)

        # --- results table
        mid = ttk.Frame(self.root)
        mid.pack(fill="both", expand=True, **pad)
        cols = ("file", "folder", "modified")
        self.tree = ttk.Treeview(mid, columns=cols, show="headings", selectmode="extended", height=10)
        for col, text, width in (("file", "File", 320), ("folder", "Folder", 340), ("modified", "Modified", 140)):
            self.tree.heading(col, text=text)
            self.tree.column(col, width=width, anchor="w")
        scroll = ttk.Scrollbar(mid, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scroll.pack(side="left", fill="y")
        self.tree.bind("<<TreeviewSelect>>", lambda _e: self.on_select())
        self.tree.bind("<Control-a>", lambda _e: (self.select_all(), "break")[1])
        self.tree.bind("<Control-A>", lambda _e: (self.select_all(), "break")[1])
        self.tree.bind("<Double-1>", lambda _e: self.open_file())
        self.root.bind("<Control-p>", lambda _e: self.do_print())
        self.root.bind("<Control-Shift-P>", lambda _e: self.do_print_many())

    # ---- helpers
    def set_status(self, text):
        self.status.set(text)

    def selected(self):
        sel = self.tree.selection()
        return self.results[int(sel[0])] if sel else None

    def selected_items(self):
        """All selected rows, in list order."""
        return [self.results[int(i)] for i in sorted(self.tree.selection(), key=int)]

    def select_all(self):
        self.tree.focus_set()
        self.tree.selection_set(self.tree.get_children())

    def run_bg(self, work, start_msg):
        if self.busy:
            return
        self.busy = True
        self.set_status(start_msg)

        def target():
            try:
                self.q.put(("done", work()))
            except Exception as exc:  # noqa: BLE001
                self.q.put(("error", str(exc)))

        threading.Thread(target=target, daemon=True).start()

    def progress(self, text):
        self.q.put(("status", text))

    def _poll(self):
        try:
            while True:
                kind, payload = self.q.get_nowait()
                if kind == "results":
                    self.show_results(*payload)
                elif kind == "plan":
                    key, plan = payload
                    self.plan[key] = plan
                    cur = self.selected()
                    if cur and (cur["path"], cur.get("variant")) == key:
                        self.show_plan(plan, cur.get("variant"))
                elif kind == "status":
                    self.set_status(payload)
                elif kind == "watch":
                    self.watch_status.set(payload)
                elif kind == "batch_found":
                    self.fill_batch(payload)
                elif kind == "done":
                    self.busy = False
                    if isinstance(payload, dict):
                        self.set_status(payload.get("msg", "Done"))
                        self.show_pdf_result(payload)
                    elif isinstance(payload, list):
                        self.set_status("Done: " + ", ".join(payload))
                    else:
                        self.set_status(str(payload))
                elif kind == "error":
                    self.busy = False
                    self.set_status("Error")
                    messagebox.showerror("Invoice Printer", payload)
        except queue.Empty:
            pass
        self.root.after(100, self._poll)

    # ---- printer
    def current_printer(self):
        return self.cfg.get("printer", "")

    def on_printer(self):
        value = self.printer_var.get()
        self.cfg["printer"] = "" if value == DEFAULT_PRINTER_LABEL else value
        save_config(self.cfg)

    def refresh_printers(self):
        printers = [DEFAULT_PRINTER_LABEL] + list_printers()
        self.printer_box["values"] = printers
        if self.printer_var.get() not in printers:
            self.printer_var.set(DEFAULT_PRINTER_LABEL)
            self.on_printer()
        self.set_status(f"{len(printers) - 1} printer(s) found.")

    def open_printer_setup(self):
        """Open the Windows 'Printing Preferences' of the chosen printer. Save the printer's department code /
        user PIN / password THERE once, and every print job from this app carries it (no prompt each time)."""
        name = self.current_printer()
        if not name:
            try:
                import win32print
                name = win32print.GetDefaultPrinter()
            except Exception:
                name = ""
        if not name:
            messagebox.showinfo("Printer setup", "No printer selected.")
            return
        try:
            import subprocess
            subprocess.Popen(["rundll32", "printui.dll,PrintUIEntry", "/e", "/n", name])
            self.set_status(f"Printing Preferences opened for {name} - save the department code / PIN there.")
        except Exception as exc:
            messagebox.showerror("Printer setup", f"Could not open printer preferences:\n{exc}")

    def printer_label(self):
        return self.current_printer() or "the Windows default printer"

    # ---- folders
    def browse_root(self):
        folder = filedialog.askdirectory(title="Choose the top folder that contains the month folders",
                                         initialdir=self.cfg["root_folder"] or None)
        if folder:
            self.cfg["root_folder"] = folder
            self.root_var.set(folder)
            if not self.cfg["watch_folder"]:
                self.cfg["watch_folder"] = folder
                self.watch_var.set(folder)
            save_config(self.cfg)

    def browse_watch(self):
        folder = filedialog.askdirectory(title="Choose the folder to watch for NEW invoice files",
                                         initialdir=self.cfg["watch_folder"] or None)
        if not folder:
            return
        self.cfg["watch_folder"] = folder
        self.watch_var.set(folder)
        save_config(self.cfg)
        if self.auto_var.get():          # restart the watcher on the new folder
            self.stop_watch()
            self.start_watch()

    # ---- auto-print
    def toggle_auto(self):
        if self.auto_var.get():
            folder = self.cfg["watch_folder"]
            if not folder or not os.path.isdir(folder):
                messagebox.showwarning("Auto-print", "Choose a valid folder to watch first.")
                self.auto_var.set(False)
                return
            ok = messagebox.askyesno(
                "Turn on auto-print?",
                f"Every NEW Excel file that appears in:\n{folder}\n\n"
                f"will be printed on {self.printer_label()} ({int(self.copies_var.get())} copy/copies) "
                f"without asking.\n\nFiles that are already there are ignored. Continue?")
            if not ok:
                self.auto_var.set(False)
                return
            self.cfg["copies"] = max(1, int(self.copies_var.get()))
            self.cfg["auto_print_enabled"] = True
            save_config(self.cfg)
            self.start_watch()
        else:
            self.cfg["auto_print_enabled"] = False
            save_config(self.cfg)
            self.stop_watch()

    def start_watch(self):
        self.stop_watch()
        self.watch_stop = threading.Event()
        stop = self.watch_stop
        threading.Thread(target=watch_loop, args=(stop, self.cfg, lambda m: self.q.put(("watch", m))),
                         daemon=True).start()
        self.watch_status.set("Starting…")

    def stop_watch(self):
        if self.watch_stop:
            self.watch_stop.set()
            self.watch_stop = None
        self.watch_status.set("Off")

    # ---- single search / print
    def get_query(self):
        """Text box -> one clean query. Lines are joined with commas, so pasted lists just work."""
        lines = [ln.strip() for ln in self.entry.get("1.0", "end").splitlines() if ln.strip()]
        return ", ".join(lines)

    # ---- PDF mode helpers
    def pdf_mode(self):
        return bool(self.pdf_var.get())

    def search_exts(self):
        return self.cfg.get("pdf_extensions", [".pdf"]) if self.pdf_mode() else self.cfg["extensions"]

    def search_target(self):
        return self.cfg.get("pdf_target_folder_name", "") if self.pdf_mode() else self.cfg["target_folder_name"]

    def print_btn_text(self):
        return "Print PDF" if self.pdf_mode() else "Print INV + PL"

    def on_pdf_toggle(self):
        self.cfg["print_pdf_files"] = self.pdf_mode()
        save_config(self.cfg)
        self.btn_print.configure(text=self.print_btn_text())
        try:
            self.all_chk.configure(state="disabled" if self.pdf_mode() else "normal")
        except Exception:
            pass
        self.results = []
        self.plan = {}
        self.tree.delete(*self.tree.get_children())
        self.plan_var.set("")
        self.set_status("PDF mode: searching .pdf files and printing them as they are." if self.pdf_mode()
                        else "Excel mode: printing the INV and PL sheets.")
        if self.get_query():
            self.search()

    def on_all_toggle(self):
        self.cfg["all_variants"] = bool(self.all_var.get())
        save_config(self.cfg)
        if self.get_query():
            self.search()

    def update_count(self):
        n = len(parse_numbers(self.get_query()))
        self.count_var.set(f"{n} invoice numbers entered" if n > 1 else "")

    def import_excel(self):
        path = filedialog.askopenfilename(
            title="Choose Excel / CSV with invoice numbers",
            filetypes=[("Excel / CSV", "*.xlsx *.xlsm *.csv"), ("All files", "*.*")])
        if not path:
            return
        try:
            numbers = read_numbers_from_file(path)
        except Exception as exc:  # noqa: BLE001
            return messagebox.showerror("Import failed", str(exc))
        if not numbers:
            if messagebox.askyesno("No invoice numbers found",
                                   "No invoice numbers were found in the file.\n\nExpected: a column headed "
                                   "“Invoice no.” (optional column “Variant” with 1, 2 …), one invoice per row.\n\n"
                                   "Save an example template?"):
                out = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="invoice_list_template.xlsx",
                                                   filetypes=[("Excel", "*.xlsx")], title="Save template as")
                if out:
                    save_template(out)
                    self.set_status("Template saved: " + out)
            return
        self.entry.delete("1.0", "end")
        self.entry.insert("1.0", "\n".join(numbers))
        self.update_count()
        self.set_status(f"Imported {len(numbers)} invoice number(s) from {os.path.basename(path)}")
        self.search()

    def clear_query(self):
        self.entry.delete("1.0", "end")
        self.count_var.set("")
        self.entry.focus_set()

    def search(self):
        query = self.get_query()
        pdf = self.pdf_mode()
        exts, target = self.search_exts(), self.search_target()
        if query and self.all_var.get() and not pdf:          # 'C264958' -> 'C264958-*'  (every INV / PL sheet)
            query = ", ".join(with_all_variants(parse_numbers(query)))
        if not self.cfg["root_folder"]:
            return self.browse_root()
        if not query:
            return
        self.cancel.set()  # stop any older search
        self.cancel = threading.Event()
        cancel = self.cancel
        self.set_status(f"Searching for “{query}”…")

        multi = None if pdf else parse_multi(query)   # 'C264958,-1,-2' -> one file, sheets INV/INV-1/INV-2 (+ PL)
        search_q = multi[0] if multi else query
        numbers = parse_numbers(query)

        if not multi and len(numbers) > 1:   # 'C264958, C264959, C264960' -> find ALL at once
            self.set_status(f"Searching {len(numbers)} invoice numbers…" + (" (PDF)" if pdf else ""))

            def work_many():
                found = find_many(self.cfg["root_folder"], numbers, target, exts, cancel,
                                  base_fallback=not pdf)
                rows = []
                for n in numbers:
                    if found[n]:
                        rows.append(dict(found[n][0], number=n, multi=True, label=f"{n}   →   {found[n][0]['name']}",
                                         **({"variant": None} if pdf else {})))
                    else:
                        rows.append({"path": None, "name": n, "folder": "", "mtime": 0, "variant": None,
                                     "number": n, "multi": True, "label": f"{n}   →   ✖ not found"})
                self.q.put(("results", (rows, False)))

            threading.Thread(target=work_many, daemon=True).start()
            return

        def work():
            hits = find_invoices(self.cfg["root_folder"], search_q, target, exts, cancel, base_fallback=not pdf)
            if pdf:
                for h in hits:
                    h["variant"] = None
            fallback = False
            if not hits and target and not cancel.is_set():
                hits = find_invoices(self.cfg["root_folder"], search_q, "", exts, cancel, base_fallback=not pdf)
                fallback = bool(hits)
            if multi:
                for h in hits:
                    h["variant"] = multi[1]
            self.q.put(("results", (hits, fallback)))

        threading.Thread(target=work, daemon=True).start()

    def show_results(self, hits, fallback=False):
        self.results = hits
        self.tree.delete(*self.tree.get_children())
        self.plan_var.set("")
        for i, h in enumerate(hits):
            stamp = datetime.fromtimestamp(h["mtime"]).strftime("%Y-%m-%d %H:%M") if h["mtime"] else ""
            self.tree.insert("", "end", iid=str(i), values=(h.get("label", h["name"]), h["folder"], stamp))
        n_found = sum(1 for h in hits if h.get("path"))
        self.btn_print_all.configure(state="normal" if n_found else "disabled",
                                    text=f"Print ALL found ({n_found})" if n_found else "Print ALL found")
        if not self.pdf_mode():
            self._prefetch_plans(hits)
        if not hits:
            self.set_status("No invoice found under: " + self.cfg["root_folder"])
            return
        if hits[0].get("multi"):
            ok = sum(1 for h in hits if h["path"])
            miss = [h["number"] for h in hits if not h["path"]]
            self.set_status(f"{ok} of {len(hits)} found – “Print ALL found” prints every found invoice"
                            + (f"  (not found: {', '.join(miss)})" if miss else "") + ".")
            self.tree.selection_set("0")
            self.tree.focus("0")
            return
        note = " (none inside an inv&pl folder, showing matches from all folders)" if fallback else ""
        self.set_status(f"{len(hits)} file(s) found{note}.  “{self.print_btn_text()}” = selected row,"
                        f"  “Print ALL found” = all {len(hits)}.")
        self.tree.selection_set("0")
        self.tree.focus("0")

    def _prefetch_plans(self, hits):
        """Look at the sheets of every found file in the background so 'Print ALL' can list them."""
        todo = [(h["path"], h.get("variant")) for h in hits
                if h.get("path") and (h["path"], h.get("variant")) not in self.plan]

        def work():
            for key in todo:
                try:
                    plan = plan_sheets_from_file(key[0], self.cfg, key[1])
                except Exception:
                    plan = None
                self.q.put(("plan", (key, plan)))

        if todo:
            threading.Thread(target=work, daemon=True).start()

    def on_select(self):
        items = self.selected_items()
        found = [h for h in items if h.get("path")]
        if self.pdf_mode():
            if len(items) > 1:
                self.btn_print.configure(text=f"Print selected ({len(found)})")
                self.plan_var.set(f"{len(items)} PDF rows selected ({len(found)} found) – every PDF is printed as it is.")
            else:
                self.btn_print.configure(text="Print PDF")
                item = self.selected()
                if item:
                    self.plan_var.set(("PDF – the whole file will be printed: " + item["name"]) if item.get("path")
                                      else "⚠ File not found for " + item.get("number", ""))
            return
        if len(items) > 1:
            self.btn_print.configure(text=f"Print selected ({len(found)})")
            self.plan_var.set(f"{len(items)} rows selected ({len(found)} found) – “Print selected” prints the "
                              "INV + PL sheets of every selected invoice.")
            return
        self.btn_print.configure(text=self.print_btn_text())
        item = self.selected()
        if not item:
            return
        path = item["path"]
        if not path:
            return self.plan_var.set("⚠ File not found for " + item.get("number", ""))
        key = (path, item.get("variant"))
        if key in self.plan:
            return self.show_plan(self.plan[key], item.get("variant"))
        self.plan_var.set("Checking sheets…")

        def work():
            try:
                plan = plan_sheets_from_file(path, self.cfg, key[1])
            except Exception:
                plan = None
            self.q.put(("plan", (key, plan)))

        threading.Thread(target=work, daemon=True).start()

    def show_plan(self, plan, variant=None):
        if plan is None:
            self.plan_var.set("Sheets will be detected when printing.")
            return
        names, picked, empty, dupes = plan
        if picked:
            dup_names = [a for a, _b in dupes]
            skipped = [n for n in names if n not in picked and n not in empty and n not in dup_names]
            text = "Will print: " + "  •  ".join(picked)
            if empty:
                text += f"      (empty, not printed: {', '.join(empty)})"
            if dupes:
                text += "      (duplicate, not printed: " + ", ".join(f"{a} = {b}" for a, b in dupes) + ")"
            if skipped:
                text += f"      (skipped: {', '.join(skipped)})"
            warn = pair_warning(picked)
            if warn:
                text += f"      ⚠ {warn}"
            self.plan_var.set(text)
        else:
            msg = no_sheet_message(variant) if variant is not None else "No INV/PL sheet found."
            self.plan_var.set("⚠ " + msg + " Sheets: " + ", ".join(names))

    def do_print_many(self, only=None):
        """Print INV + PL of every found row - all rows, or only `only` (the selected rows)."""
        if self.busy:
            return
        source = only if only is not None else self.results
        rows = [h for h in source if h.get("path")]
        missing = [h.get("number", h["name"]) for h in source if not h.get("path")]
        if not rows:
            return messagebox.showwarning("Nothing to print", "No invoice file was found.")
        copies = max(1, int(self.copies_var.get()))
        self.cfg["copies"] = copies
        save_config(self.cfg)

        def line(h):
            plan = self.plan.get((h["path"], h.get("variant")))
            sheets = ("  [" + ", ".join(plan[1]) + "]") if plan and plan[1] else ""
            if plan is not None and not plan[1]:
                sheets = "  [⚠ no INV/PL sheet]"
            return f"• {h.get('number', '')}  →  {h['name']}{sheets}".replace("•   →", "•")

        what_all = "selected" if only is not None else "ALL"
        if self.pdf_mode():
            def line(h):  # noqa: F811
                return f"• {h.get('number', '')}  →  {h['name']}".replace("•   →", "•")
        msg = (f"Print {what_all} {len(rows)} {'PDF file' if self.pdf_mode() else 'invoice'}(s) on {self.printer_label()}?"
               f"\nCopies each: {copies}\n\n" + "\n".join(line(h) for h in rows))
        if missing:
            msg += "\n\nNot found, skipped: " + ", ".join(missing)
        if not messagebox.askyesno("Print selected?" if only is not None else "Print all?", msg):
            return
        printer = self.current_printer()
        paths = [h["path"] for h in rows]
        variants = [h.get("variant") for h in rows]

        pdf = self.pdf_mode()            # read the checkbox here (GUI thread), not inside the worker thread

        def work():
            if pdf:
                res = print_pdf_files(paths, self.cfg, copies, printer, "print-all", self.progress)
            else:
                res = process_files(paths, self.cfg, "print", copies=copies, printer=printer,
                                    source="print-all", progress=self.progress, variants=variants)
            bad = [r for r in res if not r["ok"]]
            if bad:
                raise RuntimeError(f"{len(res) - len(bad)} printed, {len(bad)} failed:\n\n"
                                   + "\n".join(f"• {r['name']}: {r['error']}" for r in bad)
                                   + "\n\nDetails are in the print log.")
            return f"Printed {len(res)} invoice(s)."

        self.run_bg(work, f"Printing {len(paths)} invoice(s)…")

    def do_print(self):
        if self.busy:
            return
        chosen = self.selected_items()
        if len(chosen) > 1:
            return self.do_print_many(only=chosen)
        item = self.selected()
        if not item:
            return
        if not item.get("path"):
            return messagebox.showwarning("Not found", "Select a row whose file was found "
                                                       "(or use “Print ALL found”).")
        if self.pdf_mode():
            return self.print_one_pdf(item)
        known = self.plan.get((item["path"], item.get("variant")))
        if known is not None and not known[1]:
            messagebox.showwarning("Nothing to print", no_sheet_message(item.get("variant")))
            return
        copies = max(1, int(self.copies_var.get()))
        self.cfg["copies"] = copies
        save_config(self.cfg)
        if self.cfg.get("confirm_before_print", True):
            plan = self.plan.get((item["path"], item.get("variant")))
            what = "  •  ".join(plan[1]) if plan else "the INV and PL sheets"
            if not messagebox.askyesno("Print?", f"{item['name']}\n\n{what}\n\n"
                                                 f"Printer: {self.printer_label()}\nCopies: {copies}"):
                return
        printer = self.current_printer()

        def work():
            res = process_files([item["path"]], self.cfg, "print", copies=copies, printer=printer,
                                source="manual", progress=self.progress,
                                variants=[item.get("variant")])[0]
            if not res["ok"]:
                raise RuntimeError(res["error"])
            return res["sheets"]

        self.run_bg(work, f"Printing {item['name']}…")

    def print_one_pdf(self, item):
        copies = max(1, int(self.copies_var.get()))
        self.cfg["copies"] = copies
        save_config(self.cfg)
        if self.cfg.get("confirm_before_print", True):
            if not messagebox.askyesno("Print PDF?", f"{item['name']}\n\nThe whole PDF file will be printed.\n\n"
                                                      f"Printer: {self.printer_label()}\nCopies: {copies}"):
                return
        printer = self.current_printer()

        def work():
            res = print_pdf_files([item["path"]], self.cfg, copies, printer, "manual", self.progress)[0]
            if not res["ok"]:
                raise RuntimeError(res["error"])
            return res["sheets"]

        self.run_bg(work, f"Printing {item['name']}…")

    def do_pdf(self):
        item = self.selected()
        if not item or self.busy:
            return
        if self.pdf_mode():
            return messagebox.showinfo("Already PDF", "In PDF mode the files are already PDF.\n\n"
                                       "Use “Print PDF”, or “Combine → 1 PDF” to merge the selected files into one.")
        if len(self.tree.selection()) > 1:
            return self.do_combine_pdf()
        if not item.get("path"):
            return messagebox.showwarning("Not found", "Select a row whose file was found.")
        known = self.plan.get((item["path"], item.get("variant")))
        if known is not None and not known[1]:
            messagebox.showwarning("Nothing to print", no_sheet_message(item.get("variant")))
            return
        default = Path(item["name"]).stem + "_INV-PL.pdf"
        out = filedialog.asksaveasfilename(defaultextension=".pdf", initialfile=default,
                                           filetypes=[("PDF", "*.pdf")], title="Save PDF as")
        if not out:
            return

        def work():
            res = process_files([item["path"]], self.cfg, "pdf", pdf_path=out, source="manual",
                                progress=self.progress, variants=[item.get("variant")])[0]
            if not res["ok"]:
                raise RuntimeError(res["error"])
            if not os.path.exists(out):
                raise RuntimeError("Excel reported success but no PDF file was produced:\n" + out)
            return {"msg": "PDF saved: " + ", ".join(res["sheets"]), "path": out}

        self.run_bg(work, f"Creating PDF for {item['name']}…")

    def show_pdf_result(self, info):
        """Tell the user where the PDF is and open it (runs on the GUI thread)."""
        import subprocess
        out = info.get("path", "")
        if not out or not os.path.exists(out):
            return messagebox.showerror("PDF not created", info.get("msg", "") + "\n\nThe PDF file was not found:\n" + out)
        opened = True
        try:
            os.startfile(out)
        except Exception:
            opened = False
            try:
                subprocess.Popen(["explorer", "/select,", os.path.normpath(out)])
            except Exception:
                pass
        size_kb = os.path.getsize(out) // 1024
        messagebox.showinfo("PDF ready", f"{info.get('msg', '')}\n\nFile: {out}  ({size_kb} KB)"
                            + ("" if opened else "\n\nCould not open a PDF viewer - the folder was opened instead."))

    def do_combine_pdf(self):
        """Selected rows (or ALL found rows if none selected) -> one combined PDF, ready to print once."""
        if self.busy:
            return
        rows = self.selected_items() or list(self.results)
        rows = [h for h in rows if h.get("path")]
        if not rows:
            return messagebox.showwarning("Nothing to combine", "No invoice file was found.")
        default = datetime.now().strftime("Combined_%Y%m%d_%H%M") + ".pdf"
        out = filedialog.asksaveasfilename(defaultextension=".pdf", initialfile=default,
                                           filetypes=[("PDF", "*.pdf")],
                                           title=f"Save combined PDF ({len(rows)} invoice(s))")
        if not out:
            return
        cfg = self.cfg
        if self.pdf_mode():
            def work_pdf():
                pages = merge_pdfs([r["path"] for r in rows], out, [Path(r["name"]).stem for r in rows])
                return {"msg": f"Combined {len(rows)} PDF file(s) → {pages} page(s).", "path": out}

            return self.run_bg(work_pdf, f"Combining {len(rows)} PDF file(s)…")

        def work():
            pages, ok, bad = export_and_combine(rows, cfg, out, progress=self.progress)
            msg = f"Combined {ok} of {len(rows)} invoice(s) → {pages} page(s)."
            if bad:
                msg += f"\n\n{len(bad)} failed:\n" + "\n".join(f"• {r['name']}: {r['error']}" for r in bad)
            return {"msg": msg, "path": out}

        self.run_bg(work, f"Combining {len(rows)} invoice(s) into one PDF…")

    def open_file(self):
        item = self.selected()
        if item and item.get("path"):
            os.startfile(item["path"])  # Windows only

    # ---- batch printing
    def open_batch(self):
        if not self.cfg["root_folder"]:
            return self.browse_root()
        if self.batch_win is not None and self.batch_win.winfo_exists():
            self.batch_win.lift()
            return
        win = tk.Toplevel(self.root)
        self.batch_win = win
        win.title("Batch print")
        win.geometry("860x600")
        win.transient(self.root)
        frm = ttk.Frame(win, padding=14)
        frm.pack(fill="both", expand=True)

        ttk.Label(frm, text="Paste invoice numbers (one per line, or separated by commas / spaces):").pack(anchor="w")
        self.batch_text = tk.Text(frm, height=6, font=("Consolas", 11))
        self.batch_text.pack(fill="x", pady=(4, 8))
        self.batch_text.focus_set()

        row = ttk.Frame(frm)
        row.pack(fill="x")
        ttk.Button(row, text="Find files", style="Primary.TButton", command=self.batch_search).pack(side="left")
        self.batch_info = tk.StringVar(value="")
        ttk.Label(row, textvariable=self.batch_info, style="Hint.TLabel").pack(side="left", padx=12)

        cols = ("number", "state", "file", "folder")
        self.batch_tree = ttk.Treeview(frm, columns=cols, show="headings", height=12)
        for col, text, width in (("number", "Invoice no.", 130), ("state", "Result", 170),
                                 ("file", "File that will be printed", 300), ("folder", "Folder", 220)):
            self.batch_tree.heading(col, text=text)
            self.batch_tree.column(col, width=width, anchor="w")
        self.batch_tree.tag_configure("missing", foreground="#b00020")
        self.batch_tree.tag_configure("multi", foreground="#a65e00")
        self.batch_tree.pack(fill="both", expand=True, pady=(10, 8))

        bottom = ttk.Frame(frm)
        bottom.pack(fill="x")
        self.btn_batch_print = ttk.Button(bottom, text="Print all found", style="Primary.TButton",
                                          command=self.batch_print, state="disabled")
        self.btn_batch_print.pack(side="right")
        ttk.Label(bottom, text="If several files match a number, the newest one inside an inv&pl folder is used.",
                  style="Hint.TLabel").pack(side="left")

    def batch_search(self):
        numbers = parse_numbers(self.batch_text.get("1.0", "end"))
        if not numbers:
            return
        self.batch_info.set(f"Searching {len(numbers)} number(s)…")
        self.btn_batch_print.configure(state="disabled")
        self.batch_tree.delete(*self.batch_tree.get_children())
        cfg = self.cfg
        pdf, target, exts = self.pdf_mode(), self.search_target(), self.search_exts()   # GUI thread

        def work():
            found = find_many(cfg["root_folder"], numbers, target, exts, base_fallback=not pdf)
            rows = []
            for n in numbers:
                hits = found[n]
                if not hits:
                    rows.append({"number": n, "state": "Not found", "hit": None, "tag": "missing"})
                elif len(hits) == 1:
                    rows.append({"number": n, "state": "Found", "hit": hits[0], "tag": ""})
                else:
                    rows.append({"number": n, "state": f"{len(hits)} matches – newest used",
                                 "hit": hits[0], "tag": "multi"})
            self.q.put(("batch_found", rows))

        threading.Thread(target=work, daemon=True).start()

    def fill_batch(self, rows):
        if self.batch_win is None or not self.batch_win.winfo_exists():
            return
        self.batch_rows = rows
        self.batch_tree.delete(*self.batch_tree.get_children())
        for r in rows:
            hit = r["hit"]
            self.batch_tree.insert("", "end", tags=(r["tag"],) if r["tag"] else (),
                                   values=(r["number"], r["state"], hit["name"] if hit else "—",
                                           hit["folder"] if hit else ""))
        ok = sum(1 for r in rows if r["hit"])
        self.batch_info.set(f"{ok} of {len(rows)} found.")
        self.btn_batch_print.configure(state="normal" if ok else "disabled")

    def batch_print(self):
        if self.busy:
            messagebox.showinfo("Batch print", "Another print job is still running.")
            return
        rows = [r for r in self.batch_rows if r["hit"]]
        if not rows:
            return
        copies = max(1, int(self.copies_var.get()))
        missing = len(self.batch_rows) - len(rows)
        msg = f"Print {len(rows)} invoice(s) on {self.printer_label()}?\nCopies each: {copies}"
        if missing:
            msg += f"\n\n{missing} number(s) were not found and will be skipped."
        if not messagebox.askyesno("Batch print", msg, parent=self.batch_win):
            return
        self.cfg["copies"] = copies
        save_config(self.cfg)
        paths = [r["hit"]["path"] for r in rows]
        variants = [None if self.pdf_mode() else r["hit"].get("variant") for r in rows]  # GUI thread
        printer = self.current_printer()

        pdf = self.pdf_mode()            # GUI thread

        def work():
            if pdf:
                res = print_pdf_files(paths, self.cfg, copies, printer, "batch", self.progress)
            else:
                res = process_files(paths, self.cfg, "print", copies=copies, printer=printer,
                                    source="batch", progress=self.progress, variants=variants)
            ok = [r for r in res if r["ok"]]
            bad = [r for r in res if not r["ok"]]
            if bad:
                raise RuntimeError(f"{len(ok)} printed, {len(bad)} failed:\n\n"
                                   + "\n".join(f"• {r['name']}: {r['error']}" for r in bad)
                                   + "\n\nDetails are in the print log.")
            return f"Batch done: {len(ok)} invoice(s) printed."

        self.run_bg(work, f"Batch printing {len(paths)} file(s)…")
        self.batch_win.destroy()

    # ---- print log window
    def open_log(self):
        win = tk.Toplevel(self.root)
        win.title("Print log")
        win.geometry("1100x480")
        win.transient(self.root)
        cols = LOG_FIELDS
        tree = ttk.Treeview(win, columns=cols, show="headings")
        widths = {"time": 140, "mode": 90, "file": 230, "sheets": 120, "copies": 55,
                  "printer": 150, "status": 70, "message": 300}
        for c in cols:
            tree.heading(c, text=c.capitalize())
            tree.column(c, width=widths[c], anchor="w")
        tree.tag_configure("fail", foreground="#b00020")
        tree.tag_configure("skip", foreground="#a65e00")
        sb = ttk.Scrollbar(win, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=sb.set)

        def refresh():
            tree.delete(*tree.get_children())
            for r in read_log():
                r = (r + [""] * len(cols))[:len(cols)]
                status = r[6]
                tag = ("fail",) if status == "FAILED" else ("skip",) if status == "SKIPPED" else ()
                tree.insert("", "end", values=r, tags=tag)

        def row_of(item):
            return [str(v) for v in tree.item(item, "values")]

        def do_delete_selected():
            sel = tree.selection()
            if not sel:
                messagebox.showinfo("Print log", "Select one or more rows first (Ctrl/Shift + click).", parent=win)
                return
            if not messagebox.askyesno("Print log", f"Delete {len(sel)} selected row(s) from the log?", parent=win):
                return
            delete_log_rows([row_of(i) for i in sel])
            refresh()

        def do_clear_old():
            if not messagebox.askyesno(
                    "Print log", "Delete all log rows older than 7 days?\n(A backup is kept in print_log_backup.csv)",
                    parent=win):
                return
            n = clear_log(older_than_days=7)
            refresh()
            messagebox.showinfo("Print log", f"{n} old row(s) removed.", parent=win)

        def do_clear_all():
            if not messagebox.askyesno(
                    "Print log", "Clear the ENTIRE print log?\n(A backup is kept in print_log_backup.csv)",
                    icon="warning", parent=win):
                return
            n = clear_log()
            refresh()
            messagebox.showinfo("Print log", f"Log cleared ({n} row(s) removed).", parent=win)

        tree.configure(selectmode="extended")
        btns = ttk.Frame(win, padding=8)
        btns.pack(side="bottom", fill="x")
        ttk.Button(btns, text="Refresh", command=refresh).pack(side="left")
        ttk.Button(btns, text="Delete selected", command=do_delete_selected).pack(side="left", padx=(8, 0))
        ttk.Button(btns, text="Clear older than 7 days", command=do_clear_old).pack(side="left", padx=(8, 0))
        ttk.Button(btns, text="Clear all", command=do_clear_all).pack(side="left", padx=(8, 0))
        ttk.Button(btns, text="Open CSV in Excel",
                   command=lambda: LOG_PATH.exists() and os.startfile(str(LOG_PATH))).pack(side="left", padx=8)
        ttk.Label(btns, text=str(LOG_PATH), style="Hint.TLabel").pack(side="right")
        tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="left", fill="y")
        refresh()

    # ---- settings
    def open_settings(self):
        win = tk.Toplevel(self.root)
        win.title("Settings")
        win.transient(self.root)
        win.grab_set()
        frm = ttk.Frame(win, padding=16)
        frm.pack(fill="both", expand=True)
        fields = [
            ("Search only in folders named", "target_folder_name", self.cfg["target_folder_name"]),
            ("Sheets to print (regex, comma)", "sheet_include", ", ".join(self.cfg["sheet_include"])),
            ("Sheets to never print (regex, comma)", "sheet_exclude", ", ".join(self.cfg["sheet_exclude"])),
            ("Auto-print only file names matching (regex)", "watch_name_regex", self.cfg["watch_name_regex"]),
            ("Print only if it has data (regex, comma)", "conditional_sheets",
             ", ".join(self.cfg.get("conditional_sheets", []))),
        ]
        vars_ = {}
        for r, (label, key, value) in enumerate(fields):
            ttk.Label(frm, text=label).grid(row=r, column=0, sticky="w", pady=4)
            vars_[key] = tk.StringVar(value=value)
            ttk.Entry(frm, textvariable=vars_[key], width=40).grid(row=r, column=1, padx=(10, 0), pady=4)
        n = len(fields)
        confirm = tk.BooleanVar(value=self.cfg.get("confirm_before_print", True))
        ttk.Checkbutton(frm, text="Ask before printing (manual print)", variable=confirm).grid(
            row=n, column=0, columnspan=2, sticky="w", pady=6)
        nodup = tk.BooleanVar(value=self.cfg.get("skip_duplicate_sheets", True))
        ttk.Checkbutton(frm, text="Don't print duplicate sheets (INV + INV, INV + INV-1 …)",
                        variable=nodup).grid(row=n + 1, column=0, columnspan=2, sticky="w")

        nocopy = tk.BooleanVar(value=self.cfg.get("skip_numbered_copies", True))
        ttk.Checkbutton(frm, text="Skip Excel copies INV (2), INV (3), PL (2) ... (real INV-1, PL-1, INV-2 sheets still print)",
                        variable=nocopy).grid(row=n + 2, column=0, columnspan=2, sticky="w")

        def save():
            self.cfg["target_folder_name"] = vars_["target_folder_name"].get().strip()
            for key in ("sheet_include", "sheet_exclude", "conditional_sheets"):
                self.cfg[key] = [p.strip() for p in vars_[key].get().split(",") if p.strip()]
            rx = vars_["watch_name_regex"].get().strip()
            try:
                re.compile(rx)
            except re.error as exc:
                messagebox.showerror("Settings", f"Invalid regex for auto-print names:\n{exc}", parent=win)
                return
            self.cfg["watch_name_regex"] = rx
            self.cfg["confirm_before_print"] = confirm.get()
            self.cfg["skip_duplicate_sheets"] = nodup.get()
            self.cfg["skip_numbered_copies"] = nocopy.get()
            save_config(self.cfg)
            self.plan.clear()
            win.destroy()

        ttk.Button(frm, text="Save", style="Primary.TButton", command=save).grid(
            row=n + 3, column=1, sticky="e", pady=(10, 0))


def _report_crash(exc_type, exc, tb):
    import traceback
    text = "".join(traceback.format_exception(exc_type, exc, tb))
    try:
        with open(ERROR_LOG_PATH, "a", encoding="utf-8") as f:
            f.write(f"\n[{datetime.now():%Y-%m-%d %H:%M:%S}]\n{text}")
    except Exception:
        pass
    try:  # windowed .exe has no console, so show the error instead of dying silently
        from tkinter import messagebox as _mb
        _mb.showerror("Invoice Printer - unexpected error",
                      f"{exc}\n\nDetails were saved to:\n{ERROR_LOG_PATH}")
    except Exception:
        pass


def main():
    if tk is None:
        sys.exit("tkinter is not available in this Python install.")
    sys.excepthook = _report_crash
    try:  # crisp text on high-DPI screens
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    root = tk.Tk()
    root.report_callback_exception = lambda et, e, tb: _report_crash(et, e, tb)  # errors inside buttons/callbacks
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
